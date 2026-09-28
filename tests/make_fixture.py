"""Regenerate tests/data: a small synthetic workbook in the real project's layout, plus the
control file built from it (golden), a .par and a .res file for the update tests.

    python tests/make_fixture.py
"""
import os
import sys

import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, 'data')
sys.path.insert(0, os.path.dirname(HERE))

from makepst.cli import main  # noqa: E402

BOOK = os.path.join(DATA, 'demo.xlsx')
PST = os.path.join(DATA, 'demo.pst')

BUILD_ARGS = [
    PST, 'regul',
    '--set_ctl_xls', f'{BOOK},CONTROL',
    '--add_pargp_xls', f'{BOOK},PARGP',
    '--add_par_xls', f'{BOOK},PAR_HK',
    '--add_par_xls', f'{BOOK},PAR_SY',
    '--add_par_xls', f'{BOOK},PAR_RCH',
    '--add_obs_xls', f'{BOOK},OBS_HEAD',
    '--add_obs_xls', f'{BOOK},OBS_FLOW',
    '--add_io_xls', f'{BOOK},IO',
    '--add_pp_xls', f'{BOOK},PPcntl',
    '--add_comment', 'demo01 synthetic fixture',
    '--add_comment', 'demo02 second history line',
    '--no_dump_tpl',
]


def sheets():
    control = pd.DataFrame([
        # LINE, NAME, DEFAULT, VALUE  -- blanks in VALUE fall back to the package default
        (2, 'pestmode', 'estimation', 'estimation'),      # overridden by the command line 'regul'
        (2, 'rstfle', 'restart', 'restart'),
        (4, 'precis', 'single', 'double'),
        (5, 'rlambda1', 10, 20),
        (5, 'rlamfac', 2, 2),
        (5, 'jacupdate', 0, 999),
        (5, 'lamforgive', 'nolamforgive', 'lamforgive'),
        (5, 'win_mrun_hours', None, 1.5),
        (5, 'uptestmin', None, 70),
        (6, 'relparmax', 3, 3),
        (6, 'facparmax', 3, 3),
        (6, 'absparmax', None, 'absparmax(1)=0.1 absparmax(2)=20'),
        (7, 'doaui', None, 'noaui'),
        (8, 'noptmax', 1, 10),
        (8, 'phiredstp', 0.005, 0.01),
        (8, 'phistopthresh', 0, 0.5),
        (9, 'ires', None, 1),
        (9, 'jcosave', 'jcosave', 'jcosave'),
        (9, 'rrfsave', 'rrfsave', 'rrfsave'),
        (0, 'svdmode', 1, 2),                              # not in the control data section
        (0, 'maxsing', 10000, 50),
        (0, 'phimlim', 0.1, 5.0),
        (0, 'wfinit', 1.0, 0.5),
        (0, 'npar', None, 999),                            # computed; must be ignored
        (0, 'not_a_variable', None, 1),                    # unknown; must be ignored
    ], columns=['LINE', 'NAME', 'DEFAULT', 'VALUE'])

    pargp = pd.DataFrame([
        ('hk', 'relative', 0.1, 0, 'always_2', 2, 'parabolic'),
        ('sy', 'relative', 0.1, 0, 'always_2', 2, 'parabolic'),
        ('rch', 'absolute', 0.01, 0, 'always_2', 2, 'parabolic'),
        ('unused', 'relative', 0.1, 0, 'switch', 2, 'parabolic'),
    ], columns='PARGPNME INCTYP DERINC DERINCLB FORCEN DERINCMUL DERMTHD'.split())

    par_cols = 'PARNME PARTRANS PARCHGLIM PARVAL1 PARLBND PARUBND PARGP SCALE OFFSET DERCOM Layer Tieto NativeVal PRIOR WEIGHT comment'.split()
    par_hk = pd.DataFrame([
        ('HK1_cc01', 'log', 'factor', 49.73775, 1, 300, 'hk', 1, 0, 1, None, 'hk1_cc01', None, 'hk1_cc03', 2, 'prior to another'),
        ('hk1_cc02', 'tied', 'factor', 90.26792, 1, 300, 'hk', 1, 0, 1, None, 'hk1_cc01', None, 'hk1_cc01', 2, 'tied'),
        ('hk1_cc03', 'log', 'factor', 123.8187, 1, 300, 'hk', 1, 0, 1, None, 'hk1_cc03', None, 100, 1, 'numeric prior'),
        ('hk1_cc04', 'fixed', 'factor', 0.001007841, 1e-5, 0.9, 'hk', 1, 0, 1, None, 'hk1_cc04', None, 'hk1_cc03', 1, 'fixed: prior dropped'),
        ('hk1_cc05', 'tied', 'factor', 5.0, 1, 300, 'hk', 1, 0, 1, None, 'hk1_cc04', None, None, None, 'tied to fixed -> fixed'),
        ('hk1_cc06', 'log', 'factor', 7.5, 1, 300, 'hk', 1, 0, 1, None, 'hk1_cc06', None, 'hk1_cc01', 0, 'zero weight: no prior'),
        (None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, 'blank row'),
    ], columns=par_cols)
    par_sy = pd.DataFrame([
        ('sy1_cc01', 'log', 'factor', 0.15, 0.01, 0.35, 'sy', 1, 0, 1, None, 'sy1_cc01', None, 'sy1_cc02', 1, None),
        ('sy1_cc02', 'log', 'factor', 0.12, 0.01, 0.35, 'sy', 1, 0, 1, None, 'sy1_cc02', None, 0.2, 1, None),
    ], columns=par_cols)
    par_rch = pd.DataFrame([
        ('rchss', 'none', 'absolute(1)', 0.3428223, 0.08, 1, 'rch', 1, 0, 1, None, None, None, 0.35, 10, None),
        ('rch2002', 'log', 'factor', 0.87, 0.5, 1.5, 'rch', 1, 0, 1, None, None, None, 1.0, 5, None),
        ('dg00025', 'none', 'absolute(2)', 1000, 900, 1100, 'rch', 1, -1000, 1, None, None, None, None, None, 'offset'),
    ], columns=par_cols)

    obs_cols = 'OBSNME OBSVAL WEIGHT OBGNME Site used'.split()
    obs_head = pd.DataFrame([
        ('MR1158636_197611', 2542.59, 0.15, 'HEAD', 'MR1158636', None),
        ('mr1158636_197705', 2542.19, 0.15, 'head', 'MR1158636', None),
        ('mr1322063_000000', 2865.475129, 0, 'headss', 'MR1322063', None),
        ('mr1332090_000000', -999, 0, 'headss', 'MR1332090', None),
    ], columns=obs_cols)
    obs_flow = pd.DataFrame([
        ('qbs_2010', 1234567.891, 1e-6, 'qbs', None, None),
        ('lhss202211', 2727.77, 0.2, 'lhss', None, None),
    ], columns=obs_cols)

    io = pd.DataFrame([
        ('cmd', '00_RunModel.bat', 'tr'),
        ('tpl', r'PEST\pphk.tpl', r'ppl\pphk.csv'),
        ('tpl', r'PEST\ppsy.tpl', r'ppl\ppsy.csv'),
        ('tpl', r'PEST\rch.tpl', r'RCH\rch.csv'),
        ('ins', r'PEST\head.ins', r'TR\Output.head'),
        ('ins', r'PEST\flow.ins', r'TR\Output.flow'),
        (None, None, None),
    ], columns=['type', 'in', 'out'])

    pp = pd.DataFrame([
        ('overdue_resched_fac', 3.0, 'real'),
        ('overdue_giveup_fac', None, 'blank: skipped'),
        ('lambdas', '0.1,1,10,100', 'text'),
        ('max_run_fail', 2, 'int'),
    ], columns=['pp_var', 'val', 'note'])

    # a BUILD sheet as a user would write it by hand: the build command, one option per row,
    # naming this workbook by its bare file name (makepst parrep reads it)
    args = [os.path.basename(a) if a == PST else a.replace(BOOK, 'demo.xlsx') for a in BUILD_ARGS]
    rows, i = ['python makePst.py ' + ' '.join(args[:2])], 2
    while i < len(args):
        opt, val = args[i], args[i + 1] if i + 1 < len(args) and not args[i + 1].startswith('--') else None
        rows.append(f'{opt} "{val}"' if val is not None and ' ' in val else opt if val is None else f'{opt} {val}')
        i += 1 if val is None else 2
    build = pd.DataFrame({'COMMAND': rows})

    return {'CONTROL': control, 'PARGP': pargp, 'PAR_HK': par_hk, 'PAR_SY': par_sy, 'PAR_RCH': par_rch,
            'OBS_HEAD': obs_head, 'OBS_FLOW': obs_flow, 'IO': io, 'PPcntl': pp, 'BUILD': build}


def write_book():
    with pd.ExcelWriter(BOOK, engine='openpyxl') as xw:
        for name, df in sheets().items():
            df.to_excel(xw, sheet_name=name, index=False)
    # helper columns that hold formulas, as in the real workbook (update must leave them alone)
    wb = openpyxl.load_workbook(BOOK)
    for name in ('PAR_HK', 'PAR_SY', 'PAR_RCH'):
        ws = wb[name]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value:
                ws.cell(r, 11, f'=INT(MID(A{r},3,1))')     # Layer
                ws.cell(r, 13, f'=D{r}*H{r}+I{r}')          # NativeVal
    for name in ('OBS_HEAD', 'OBS_FLOW'):
        ws = wb[name]
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 6, f'=B{r}*C{r}')                    # used
    wb.save(BOOK)


def write_par_res():
    from makepst import read_pst
    pst = read_pst(PST)
    with open(os.path.join(DATA, 'demo.par'), 'w') as f:
        f.write('single point\n')
        for r in pst.par.itertuples():
            # PEST writes each parameter's own SCALE and OFFSET; fill_parval converts through them
            f.write(f'{r.PARNME:<14s} {float(r.PARVAL1) * 2:.11g}  {float(r.SCALE):.1f}  {float(r.OFFSET):.1f}\n')
    with open(os.path.join(DATA, 'demo.res'), 'w') as f:
        f.write(f'{"Name":<20s} {"Group":<10s} {"Measured":>14s} {"Modelled":>14s} {"Residual":>14s} {"Weight":>10s}\n')
        for r in pst.obs.itertuples():
            f.write(f'{r.OBSNME:<20s} {r.OBGNME:<10s} {float(r.OBSVAL):>14.6f} {float(r.OBSVAL) + 1:>14.6f} '
                    f'{-1.0:>14.6f} {float(r.WEIGHT):>10.4g}\n')

    # PESTPP-IES ensembles for iteration 3: realizations 0, 1 and base; phi says 1 is best
    par = pst.par.set_index('PARNME')['PARVAL1'].astype(float)
    ens = pd.DataFrame({'0': par * 3, '1': par * 0.5, 'base': par}).T.rename_axis('real_name')
    ens.columns = [c.upper() for c in ens.columns]                  # IES keeps the pst's spelling
    ens.to_csv(os.path.join(DATA, 'demo.3.par.csv'))
    obs = pst.obs.set_index('OBSNME')['OBSVAL'].astype(float)
    pd.DataFrame({'0': obs + 3, '1': obs + 0.5, 'base': obs + 1}).T.rename_axis('real_name') \
        .to_csv(os.path.join(DATA, 'demo.3.obs.csv'))
    phi = pd.DataFrame([(2, 6, 4.0, 1.0, 3.0, 5.0, 5.0, 3.0, 4.0),
                        (3, 9, 3.3, 1.2, 2.0, 5.0, 5.0, 2.0, 3.0)],
                       columns=['iteration', 'total_runs', 'mean', 'standard_deviation', 'min', 'max',
                                '0', '1', 'base'])
    phi.to_csv(os.path.join(DATA, 'demo.phi.actual.csv'), index=False)


if __name__ == '__main__':
    os.makedirs(DATA, exist_ok=True)
    write_book()
    main(BUILD_ARGS)
    write_par_res()
    print('fixture written to', DATA)

import json
import os
import warnings

import openpyxl
import pytest

from makepst.cli import main
from makepst.excel import load_table
from makepst.provenance import check_manifest, sha256


def test_provenance_reports_changes_and_missing(tmp_path, capsys):
    source = tmp_path / 'input.txt'
    output = tmp_path / 'output.txt'
    source.write_text('before')
    output.write_text('result')
    sidecar = tmp_path / 'output.txt.manifest.json'
    sidecar.write_text(json.dumps({'sources': [{'path': str(source), 'sha256': sha256(source)}],
                                   'output': {'path': str(output), 'sha256': sha256(output)}}))
    main(['provenance', str(output)])
    assert capsys.readouterr().out.count('unchanged') == 2
    source.write_text('after')
    output.unlink()
    with pytest.raises(SystemExit) as exc:
        main(['provenance', str(sidecar)])
    assert exc.value.code == 1
    report = capsys.readouterr().out
    assert 'changed' in report and 'missing' in report


def test_update_dry_run_does_not_write(tmp_path, capsys):
    book = tmp_path / 'book.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'PAR'
    ws.append(['PARNME', 'PARVAL1'])
    ws.append(['hk', 1.0])
    wb.save(book)
    par = tmp_path / 'values.par'
    par.write_text('single point\nhk 2.0 1.0 0.0\n')
    before = sha256(book)
    main(['update', str(book), '--par', str(par), '--dry_run', '--backend', 'openpyxl'])
    assert sha256(book) == before
    assert not os.path.exists(str(book) + '.manifest.json')
    assert 'hk' in capsys.readouterr().out


def test_build_table_warns_on_missing_formula_cache(tmp_path):
    book = tmp_path / 'formula.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'PAR'
    ws.append(['PARNME', 'PARVAL1'])
    ws.append(['hk', '=1+1'])
    wb.save(book)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        load_table(f'{book},PAR')
    assert any('no cached value' in str(w.message) for w in caught)


DATA = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data')


def _fixture_book(tmp_path):
    import shutil
    shutil.copy(os.path.join(DATA, 'demo.xlsx'), tmp_path / 'demo.xlsx')
    return tmp_path / 'demo.xlsx'


def test_build_dry_run_reports_and_writes_nothing(tmp_path, capsys):
    book = _fixture_book(tmp_path)
    with pytest.raises(SystemExit) as exc:                  # the fixture's template files are not here
        main(['build', str(book), '--dry_run'])
    assert exc.value.code == 1
    out = capsys.readouterr().out
    assert 'template file not found' in out and '(dry run: nothing written)' in out
    assert sorted(os.listdir(tmp_path)) == ['demo.xlsx']      # no control file, dump.tpl or manifest


def test_build_dry_run_from_tables_named_on_the_command_line(tmp_path, capsys):
    book = str(_fixture_book(tmp_path))
    argv = ['build', str(tmp_path / 'new.pst'), 'regul', '--set_ctl_xls', f'{book},CONTROL',
            '--add_pargp_xls', f'{book},PARGP', '--add_par_xls', f'{book},PAR_*', '--add_obs_xls', f'{book},OBS_*',
            '--add_pp_xls', f'{book},PPcntl', '--dry_run']
    with pytest.raises(SystemExit) as exc:
        main(argv)
    assert exc.value.code == 1
    out = capsys.readouterr().out
    assert 'no template files' in out and 'no model command line' in out     # left out of this command
    assert 'new.pst: 3 errors' in out and '(dry run: nothing written)' in out
    assert not (tmp_path / 'new.pst').exists() and not (tmp_path / 'new.pst.manifest.json').exists()


def test_validate_judges_a_workbook_as_build_would_write_it(tmp_path, capsys):
    """Regularisation equations come from PRIOR/WEIGHT columns when building; checking the tables before
    that reported a missing "regul" group and ungenerated prior information that build never leaves."""
    book = _fixture_book(tmp_path)
    with pytest.raises(SystemExit):
        main(['validate', str(book)])
    out = capsys.readouterr().out
    assert 'whose name starts with "regul"' not in out
    assert 'have not been generated' not in out
    assert 'Rebuild prior information' in out and '(done by build)' in out


def test_workbook_without_build_sheet_points_at_dry_run(tmp_path):
    book = tmp_path / 'plain.xlsx'
    openpyxl.Workbook().save(book)
    with pytest.raises(SystemExit) as exc:
        main(['build', str(book)])
    assert 'has no BUILD sheet' in str(exc.value) and '--dry_run' in str(exc.value)

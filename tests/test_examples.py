"""The README tutorial (examples/minimal) must keep working exactly as written."""
import os
import shutil
import sys

import openpyxl
import pandas as pd
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

from makepst import from_text, read_pst, to_text  # noqa: E402
from makepst.cli import main  # noqa: E402

MINIMAL = os.path.join(ROOT, 'examples', 'minimal')


@pytest.fixture
def example(tmp_path, monkeypatch):
    shutil.copytree(MINIMAL, tmp_path / 'minimal')
    monkeypatch.chdir(tmp_path / 'minimal')
    return tmp_path / 'minimal'


def test_tutorial(example, capsys):
    main(['build', 'model.pst', 'regul', '--set_ctl_csv', 'control.csv', '--add_pargp_csv', 'pargp.csv',
          '--add_par_csv', 'par.csv', '--add_obs_csv', 'obs.csv', '--add_io_csv', 'io.csv',
          '--add_comment', 'minimal example'])
    p = read_pst('model.pst')
    assert p.counts == {'npar': 4, 'nobs': 4, 'npargp': 2, 'nprior': 2, 'nobsgp': 3,
                        'ntplfle': 1, 'ninsfle': 2, 'pestmode': 'regularisation'}
    assert p.control['maxsing'] == 3 and p.control['lamforgive'] == 'lamforgive'
    assert p.par.set_index('PARNME').loc['hk3', 'TIETO'] == 'hk1'
    assert os.path.exists('dump.tpl')

    main(['dump', 'model.pst', 'model.xlsx'])
    assert set(pd.ExcelFile('model.xlsx').sheet_names) >= {'CONTROL', 'PARGP', 'PAR', 'OBS', 'PRIOR', 'IO', 'BUILD'}

    main(['update', 'model.xlsx', '--par', 'model.par', '--res', 'model.res', '--out', 'model-results.xlsx',
          '--backend', 'openpyxl'])
    ws = openpyxl.load_workbook('model-results.xlsx')['OBS']
    hdr = [c.value for c in ws[1]]
    assert hdr[-2:] == ['MODELLED', 'RESIDUAL'] and ws.cell(2, hdr.index('MODELLED') + 1).value == 101.45
    par = pd.read_excel('model-results.xlsx', 'PAR').set_index('PARNME')['PARVAL1']
    assert par['hk1'] == 12.5 and par['rch'] == 0.0012
    assert pd.read_excel('model.xlsx', 'PAR').set_index('PARNME')['PARVAL1']['hk1'] == 10   # original untouched

    main(['parrep', 'model.pst', 'model.par', 'model-final.pst', '--set', 'noptmax=0'])
    q = read_pst('model-final.pst')
    assert q.control['noptmax'] == 0 and q.par.set_index('PARNME').loc['hk2', 'PARVAL1'] == 18.0

    main(['parrep', 'model.xlsx', 'model.par', 'model-from-book.pst', '--set', 'noptmax=0'])
    assert to_text(read_pst('model-from-book.pst')) == to_text(q)      # workbook route == pst route


def test_raw_sections_round_trip():
    text = ('pcf\n* control data\nrestart estimation\n1 1 1 0 1\n1 1 single point\n10 2 0.3 0.03 10\n'
            '3 3 0.001\n0.1\n5 0.005 4 4 0.005 4\n1 1 1\n* parameter groups\nhk relative 0.01 0 switch 2 parabolic\n'
            '* parameter data\nhk1 log factor 10 0.1 1000 hk 1 0 1\n* observation groups\nhead\n'
            '* observation data\nh1 1.0 1.0 head\n* derivatives command line\nderiv.bat\nextderiv.dat\n'
            '* model command line\nrun.bat\n* model input/output\na.tpl a.in\nb.ins b.out\n')
    p = from_text(text)
    assert p.raw_sections['derivatives command line'] == ['deriv.bat', 'extderiv.dat']
    out = to_text(p)
    i, j = out.index('* observation data'), out.index('* model command line')
    assert '* derivatives command line\nderiv.bat\nextderiv.dat\n' in out[i:j]


def test_unknown_sections_are_preserved_verbatim_in_place():
    text = ('pcf\n* control data\nrestart estimation\n1 1 1 0 1\n1 1 single point\n10 2 0.3 0.03 10\n'
            '3 3 0.001\n0.1\n5 0.005 4 4 0.005 4\n1 1 1\n* parameter groups\n'
            'hk relative 0.01 0 switch 2 parabolic\n* parameter data\n'
            'hk1 log factor 10 0.1 1000 hk 1 0 1\n* observation groups\nhead\n'
            '* observation data\nh1 1.0 1.0 head\n'
            '* Site Specific Future Section   \n  keep leading and trailing spaces  \n\n# body comment\n'
            '++future_option(retain me here)\n'
            '* model command line\nrun.bat\n* model input/output\na.tpl a.in\nb.ins b.out\n')
    with pytest.warns(UserWarning, match=r'not understood; preserved verbatim'):
        p = from_text(text)

    name = 'site specific future section'
    assert p.raw_sections[name] == ['  keep leading and trailing spaces  ', '', '# body comment',
                                    '++future_option(retain me here)']
    assert p.raw_section_headers[name] == '* Site Specific Future Section   '
    assert p.section_order[p.raw_section_positions[name]] == name

    out = to_text(p)
    block = ('* Site Specific Future Section   \n'
             '  keep leading and trailing spaces  \n\n# body comment\n'
             '++future_option(retain me here)\n')
    assert block in out
    assert out.count('++future_option(retain me here)') == 1
    assert out.index('* observation data') < out.index(block) < out.index('* model command line')

    with pytest.warns(UserWarning, match=r'not understood; preserved verbatim'):
        q = from_text(out)
    assert q.raw_sections[name] == p.raw_sections[name]

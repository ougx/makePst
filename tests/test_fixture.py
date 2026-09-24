"""Tests on the synthetic fixture in tests/data (regenerate with `python tests/make_fixture.py`)."""
import os
import shutil
import sys

import numpy as np
import openpyxl
import pandas as pd
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

from makepst import Pst, from_text, load_table, read_pst, to_text, to_workbook, update_workbook  # noqa: E402
from makepst.cli import main  # noqa: E402
from make_fixture import BOOK, BUILD_ARGS, DATA, PST  # noqa: E402


def golden():
    with open(PST) as f:
        return f.read()


def build_args(pstfile, book=BOOK, extra=()):
    args = [a.replace(BOOK, book) for a in BUILD_ARGS]
    args[0] = str(pstfile)
    return args + list(extra)


# ---------------------------------------------------------------------- build
def test_build_reproduces_golden(tmp_path):
    out = tmp_path / 'demo.pst'
    main(build_args(out))
    assert out.read_text() == golden()


def test_golden_semantics():
    p = read_pst(PST)
    assert p.pestmode == 'regularisation'                  # command line beat the CONTROL sheet
    assert p.comments == ['demo01 synthetic fixture', 'demo02 second history line']
    par = p.par.set_index('PARNME')
    assert 'hk1_cc01' in par.index and 'HK1_cc01' not in par.index   # names lower-cased
    assert par.loc['hk1_cc02', 'PARTRANS'] == 'tied' and par.loc['hk1_cc02', 'TIETO'] == 'hk1_cc01'
    assert par.loc['hk1_cc05', 'PARTRANS'] == 'fixed'      # was tied to a fixed parameter
    assert list(p.pargp['PARGPNME']) == ['hk', 'sy', 'rch']  # 'unused' dropped
    assert p.obsgp == ['head', 'headss', 'qbs', 'lhss', 'regulhk', 'regulsy', 'regulrch']
    labels = set(p.prior['PINME'])
    assert 'hk1_cc04' not in labels and 'hk1_cc02' not in labels   # fixed / tied: no prior
    assert 'hk1_cc06' not in labels                                # zero weight: no prior
    assert 'dg00025' not in labels                                 # no PRIOR value
    eq = p.prior.set_index('PINME')['EQ']
    assert eq['hk1_cc01'] == '1.0 * log(hk1_cc01) - 1.0 * log(hk1_cc03) = 0'
    assert eq['hk1_cc03'] == '1.0 * log(hk1_cc03) = 2'
    assert eq['rchss'] == '1.0 * rchss = 0.35'
    c = p.control
    assert c['jacupdate'] == 999 and c['lamforgive'] == 'lamforgive' and c['win_mrun_hours'] == 1.5
    assert c['absparmax'] == 'absparmax(1)=0.1 absparmax(2)=20' and c['precis'] == 'double'
    assert c['svdmode'] == 2 and c['maxsing'] == 50           # SVD section read from CONTROL sheet
    assert c['phimlim'] == 5 and c['phimaccept'] == 5.25 and c['wfinit'] == 0.5
    assert c['phistopthresh'] == 0.5
    assert 'npar' not in c and 'not_a_variable' not in c
    assert p.counts == {'npar': 11, 'nobs': 6, 'npargp': 3, 'nprior': 6, 'nobsgp': 7,
                        'ntplfle': 3, 'ninsfle': 2, 'pestmode': 'regularisation'}
    assert p.pestpp == [('overdue_resched_fac', '3'), ('lambdas', '0.1,1,10,100'), ('max_run_fail', '2')]
    obs = p.obs.set_index('OBSNME')
    assert obs.loc['qbs_2010', 'OBSVAL'] == pytest.approx(1234567.891)   # full precision kept
    assert obs.loc['mr1322063_000000', 'OBSVAL'] == pytest.approx(2865.475129)


def test_golden_text_details():
    text = golden()
    assert '999        lamforgive win_mrun_hours=1.5 uptestmin=70' in text
    assert 'hk1_cc02  hk1_cc01' in text                    # tied line after the parameter table
    assert '1234567.891' in text and '2865.475129' in text


def test_read_write_identity():
    assert to_text(from_text(golden())) == golden()


def test_build_estimation_mode_has_no_prior(tmp_path):
    out = tmp_path / 'est.pst'
    args = build_args(out)
    args[1] = 'estimation'
    main(args)
    p = read_pst(out)
    assert p.pestmode == 'estimation' and p.nprior == 0
    assert '* prior information' not in out.read_text() and '* regularisation' not in out.read_text()
    assert p.obsgp == ['head', 'headss', 'qbs', 'lhss']


def test_build_steady_state_drops_storage(tmp_path):
    out = tmp_path / 'ss.pst'
    main(build_args(out, extra=['--ss']))
    p = read_pst(out)
    assert not any(n.startswith('sy') for n in p.par['PARNME'])
    assert 'sy' not in list(p.pargp['PARGPNME'])
    assert not any(l.startswith('sy') for l in p.prior['PINME']) and p.nobsgp == 6


def test_build_fill_parval(tmp_path):
    out = tmp_path / 'fill.pst'
    main(build_args(out, extra=['--fill_parval', os.path.join(DATA, 'demo.par')]))
    a = read_pst(PST).par.set_index('PARNME')['PARVAL1'].astype(float)
    b = read_pst(out).par.set_index('PARNME')['PARVAL1'].astype(float)
    assert np.allclose(b, a * 2, rtol=1e-9)


def test_build_from_csv(tmp_path):
    csvs = {}
    for sheet in ('CONTROL', 'PARGP', 'PAR_HK', 'PAR_SY', 'PAR_RCH', 'OBS_HEAD', 'OBS_FLOW', 'IO', 'PPcntl'):
        path = tmp_path / f'{sheet}.csv'
        pd.read_excel(BOOK, sheet, engine='openpyxl').to_csv(path, index=False)
        csvs[sheet] = str(path)
    args = []
    for a in build_args(tmp_path / 'csv.pst'):
        if a.startswith(BOOK + ','):
            args[-1] = args[-1].replace('_xls', '_csv')
            a = csvs[a.split(',')[1]]
        args.append(a)
    main(args)
    assert (tmp_path / 'csv.pst').read_text() == golden()


def test_missing_group_is_an_error(tmp_path):
    p = Pst('regul')
    p.add_par(load_table(f'{BOOK},PAR_HK'))
    p.add_obs(load_table(f'{BOOK},OBS_HEAD'))
    with pytest.raises(ValueError, match='without a definition'):
        p.validate()


# ---------------------------------------------------------------------- dump
def test_dump_and_rebuild(tmp_path):
    ref = read_pst(PST)
    book = tmp_path / 'dump.xlsx'
    cmd = to_workbook(ref, str(book))
    assert cmd.startswith('makepst build NEW.pst regularisation')
    args = [a.strip('"').replace('dump.xlsx', str(book)) for a in cmd.split()[2:]]
    args[0] = str(tmp_path / 'rebuilt.pst')
    main(args + ['--no_dump_tpl'])
    assert (tmp_path / 'rebuilt.pst').read_text() == golden()
    build_sheet = pd.read_excel(book, 'BUILD')
    assert ' '.join(build_sheet['COMMAND']) == cmd


def test_dump_control_sheet(tmp_path):
    book = tmp_path / 'ctl.xlsx'
    to_workbook(read_pst(PST), str(book))
    ctl = pd.read_excel(book, 'CONTROL').set_index('NAME')
    assert ctl.loc['jacupdate', 'VALUE'] == 999 and ctl.loc['pestmode', 'VALUE'] == 'regularisation'
    assert pd.isna(ctl.loc['npar', 'VALUE']) and ctl.loc['npar', 'DEFAULT'] == 11
    assert ctl.loc['svdmode', 'LINE'] == 'singular value decomposition'


def test_dump_split(tmp_path):
    book = tmp_path / 'split.xlsx'
    to_workbook(read_pst(PST), str(book), split=True)
    names = pd.ExcelFile(book).sheet_names
    assert {'PAR_HK', 'PAR_SY', 'PAR_RCH', 'OBS_HEAD', 'OBS_HEADSS', 'OBS_QBS', 'OBS_LHSS'} <= set(names)
    assert 'PAR' not in names and 'OBS' not in names
    assert list(pd.read_excel(book, 'PAR_SY')['PARNME']) == ['sy1_cc01', 'sy1_cc02']


# ---------------------------------------------------------------------- update
@pytest.fixture
def book_copy(tmp_path):
    dst = tmp_path / 'demo.xlsx'
    shutil.copy(BOOK, dst)
    return str(dst)


def _col(path, sheet, name):
    ws = openpyxl.load_workbook(path)[sheet]
    hdr = [c.value for c in ws[1]]
    i = hdr.index(name) + 1
    return [ws.cell(r, i).value for r in range(2, ws.max_row + 1)]


def test_update_par_from_parfile(book_copy, capsys):
    before = _col(book_copy, 'PAR_HK', 'PARVAL1')
    update_workbook(book_copy, par=os.path.join(DATA, 'demo.par'), backend='openpyxl')
    after = _col(book_copy, 'PAR_HK', 'PARVAL1')
    assert all(abs(b - 2 * a) < 1e-9 for a, b in zip(before, after) if a is not None)
    assert _col(book_copy, 'PAR_HK', 'NativeVal')[0] == '=D2*H2+I2'        # formula untouched
    assert _col(book_copy, 'PAR_HK', 'comment')[0] == 'prior to another'   # other columns untouched
    assert 'PAR_HK: 6 rows matched' in capsys.readouterr().out


def test_update_skips_formula_cells(book_copy, capsys):
    df = pd.DataFrame({'PARNME': ['hk1_cc01'], 'NATIVEVAL': [123.0], 'PARVAL1': [1.0]})
    update_workbook(book_copy, par=df, par_cols=('PARVAL1', 'NATIVEVAL'), backend='openpyxl')
    assert _col(book_copy, 'PAR_HK', 'NativeVal')[0] == '=D2*H2+I2'
    assert _col(book_copy, 'PAR_HK', 'PARVAL1')[0] == 1.0
    assert '1 formula cells left alone' in capsys.readouterr().out
    update_workbook(book_copy, par=df, par_cols=('NATIVEVAL',), backend='openpyxl', overwrite_formulas=True)
    assert _col(book_copy, 'PAR_HK', 'NativeVal')[0] == 123.0


def test_update_res_adds_columns(book_copy):
    update_workbook(book_copy, res=os.path.join(DATA, 'demo.res'), backend='openpyxl')
    ws = openpyxl.load_workbook(book_copy)['OBS_HEAD']
    hdr = [c.value for c in ws[1]]
    assert hdr[-2:] == ['MODELLED', 'RESIDUAL']
    assert _col(book_copy, 'OBS_HEAD', 'MODELLED')[0] == pytest.approx(2543.59)
    assert _col(book_copy, 'OBS_HEAD', 'RESIDUAL') == [-1.0] * 4
    assert _col(book_copy, 'OBS_HEAD', 'used')[0] == '=B2*C2'


def test_update_from_pst_writes_obs_and_par(book_copy, tmp_path):
    p = read_pst(PST)
    p.par['PARVAL1'] = 42.0
    p.obs['WEIGHT'] = 9.0
    out = tmp_path / 'out.xlsx'
    update_workbook(book_copy, par=p, obs=p, out=str(out), backend='openpyxl',
                    par_cols=('PARVAL1', 'PARTRANS'), obs_cols=('WEIGHT',))
    assert set(_col(out, 'PAR_SY', 'PARVAL1')) == {42.0}
    assert _col(out, 'PAR_HK', 'PARTRANS')[4] == 'fixed'      # hk1_cc05 tied->fixed written back
    assert set(_col(out, 'OBS_FLOW', 'WEIGHT')) == {9.0}
    assert set(_col(book_copy, 'OBS_FLOW', 'WEIGHT')) == {1e-6, 0.2}   # original untouched (--out)


def test_update_nothing_to_do(book_copy):
    with pytest.raises(ValueError, match='nothing to update'):
        update_workbook(book_copy, backend='openpyxl')


def test_update_sheets_filter(book_copy, capsys):
    update_workbook(book_copy, par=os.path.join(DATA, 'demo.par'), backend='openpyxl', sheets=['par_hk'])
    assert _col(book_copy, 'PAR_HK', 'PARVAL1')[0] == pytest.approx(49.73775 * 2)
    assert _col(book_copy, 'PAR_SY', 'PARVAL1')[0] == pytest.approx(0.15)        # untouched
    out = capsys.readouterr().out
    assert 'PAR_HK: 6 rows matched' in out and 'PAR_SY' not in out


def test_update_groups_filter(book_copy):
    p = read_pst(PST)
    p.par['PARVAL1'] = 42.0
    update_workbook(book_copy, par=p, backend='openpyxl', groups=['SY'])
    assert set(_col(book_copy, 'PAR_SY', 'PARVAL1')) == {42.0}
    assert _col(book_copy, 'PAR_HK', 'PARVAL1')[0] == pytest.approx(49.73775)
    update_workbook(book_copy, res=os.path.join(DATA, 'demo.res'), backend='openpyxl', groups=['head'])
    assert _col(book_copy, 'OBS_HEAD', 'RESIDUAL')[:2] == [-1.0, -1.0]        # group head
    assert _col(book_copy, 'OBS_HEAD', 'RESIDUAL')[2:] == [None, None]        # group headss
    assert 'RESIDUAL' not in [c.value for c in openpyxl.load_workbook(book_copy)['OBS_FLOW'][1]]


def test_update_groups_need_group_column(book_copy):
    with pytest.raises(ValueError, match='no PARGP column'):
        update_workbook(book_copy, par=os.path.join(DATA, 'demo.par'), backend='openpyxl', groups=['hk'])


def test_update_cli_par_with_pst_groups(book_copy):
    main(['update', book_copy, '--par', os.path.join(DATA, 'demo.par'), '--pst', PST,
          '--group', 'rch', '--sheet', 'PAR_*', '--backend', 'openpyxl'])
    assert _col(book_copy, 'PAR_RCH', 'PARVAL1')[0] == pytest.approx(0.3428223 * 2)
    assert _col(book_copy, 'PAR_HK', 'PARVAL1')[0] == pytest.approx(49.73775)
    assert _col(book_copy, 'OBS_HEAD', 'WEIGHT')[0] == pytest.approx(0.15)      # --pst obs not written: OBS_* excluded


def test_update_cli_par_with_pst_does_not_write_observations(book_copy):
    before = _col(book_copy, 'OBS_HEAD', 'WEIGHT')
    main(['update', book_copy, '--par', os.path.join(DATA, 'demo.par'), '--pst', PST,
          '--backend', 'openpyxl'])
    assert _col(book_copy, 'OBS_HEAD', 'WEIGHT') == before
    assert _col(book_copy, 'PAR_HK', 'PARVAL1')[0] == pytest.approx(49.73775 * 2)


@pytest.mark.parametrize('result_flag,result_file', [
    ('--res', 'demo.res'),
    ('--obs_csv', 'demo.3.obs.csv'),
])
def test_update_cli_observations_with_pst_preserves_parameters(book_copy, result_flag, result_file):
    wb = openpyxl.load_workbook(book_copy)
    parameter_sheets = [name for name in wb.sheetnames if name.startswith('PAR_')]
    for name in parameter_sheets:
        ws = wb[name]
        column = [c.value for c in ws[1]].index('PARVAL1') + 1
        ws.cell(2, column, 7.0)             # user edits differ from the contextual control file
        ws.cell(3, column, '=1+2')
    before = {name: list(wb[name].values) for name in parameter_sheets}
    wb.save(book_copy)
    wb.close()
    observation_weights = _col(book_copy, 'OBS_HEAD', 'WEIGHT')

    main(['update', book_copy, result_flag, os.path.join(DATA, result_file), '--pst', PST,
          '--backend', 'openpyxl'])

    after = openpyxl.load_workbook(book_copy)
    try:
        assert {name: list(after[name].values) for name in parameter_sheets} == before
        assert 'PHI' in after.sheetnames
    finally:
        after.close()
    assert _col(book_copy, 'OBS_HEAD', 'MODELLED')[0] == pytest.approx(2543.59)
    assert _col(book_copy, 'OBS_HEAD', 'RESIDUAL') == [-1.0] * 4
    assert _col(book_copy, 'OBS_HEAD', 'WEIGHT') == observation_weights


def test_update_cli_pst_only_writes_parameters_and_observations(book_copy):
    wb = openpyxl.load_workbook(book_copy)
    for sheet, field, value in [('PAR_HK', 'PARVAL1', 7.0), ('OBS_HEAD', 'WEIGHT', 8.0)]:
        ws = wb[sheet]
        column = [c.value for c in ws[1]].index(field) + 1
        ws.cell(2, column, value)
    wb.save(book_copy)
    wb.close()

    main(['update', book_copy, '--pst', PST, '--backend', 'openpyxl'])

    assert _col(book_copy, 'PAR_HK', 'PARVAL1')[0] == pytest.approx(49.73775)
    assert _col(book_copy, 'OBS_HEAD', 'WEIGHT')[0] == pytest.approx(0.15)


# ---------------------------------------------------------------------- PESTPP-IES ensembles
PAR_CSV = os.path.join(DATA, 'demo.3.par.csv')
OBS_CSV = os.path.join(DATA, 'demo.3.obs.csv')


def test_read_par_ensemble_realizations():
    from makepst import read_par
    base = read_pst(PST).par.set_index('PARNME')['PARVAL1'].astype(float)
    assert np.allclose(read_par(PAR_CSV)['PARVAL1'].reindex(base.index), base)             # default: base
    assert np.allclose(read_par(PAR_CSV, 'BASE')['PARVAL1'].reindex(base.index), base)     # case-insensitive
    assert np.allclose(read_par(PAR_CSV, 0)['PARVAL1'].reindex(base.index), base * 3)
    assert np.allclose(read_par(PAR_CSV, 'best')['PARVAL1'].reindex(base.index), base * 0.5)  # phi: 1 is best
    with pytest.raises(ValueError, match="realization '7' not in"):
        read_par(PAR_CSV, 7)


def test_ensemble_best_needs_phi_file(tmp_path):
    from makepst import read_par
    shutil.copy(PAR_CSV, tmp_path / 'demo.3.par.csv')
    with pytest.raises(FileNotFoundError, match='phi.actual.csv'):
        read_par(str(tmp_path / 'demo.3.par.csv'), 'best')
    shutil.copy(PAR_CSV, tmp_path / 'odd_name.csv')
    with pytest.raises(ValueError, match='case.N.par.csv'):
        read_par(str(tmp_path / 'odd_name.csv'), 'best')


def test_ensemble_without_base_needs_real(tmp_path):
    from makepst import read_par, read_ensemble
    ens = read_ensemble(PAR_CSV).drop(index='base')
    ens.to_csv(tmp_path / 'nobase.csv')
    with pytest.raises(ValueError, match="no 'base' realization"):
        read_par(str(tmp_path / 'nobase.csv'))
    assert len(read_par(str(tmp_path / 'nobase.csv'), '1')) == 11


def test_build_fill_parval_from_ensemble(tmp_path):
    out = tmp_path / 'ies.pst'
    main(build_args(out, extra=['--fill_parval', PAR_CSV, '--real', '0']))
    a = read_pst(PST).par.set_index('PARNME')['PARVAL1'].astype(float)
    b = read_pst(out).par.set_index('PARNME')['PARVAL1'].astype(float)
    assert np.allclose(b, a * 3)


def test_update_from_par_ensemble(book_copy):
    update_workbook(book_copy, par=PAR_CSV, real='best', backend='openpyxl')
    assert _col(book_copy, 'PAR_HK', 'PARVAL1')[0] == pytest.approx(49.73775 * 0.5)
    assert _col(book_copy, 'PAR_RCH', 'PARVAL1')[0] == pytest.approx(0.3428223 * 0.5)


def test_update_from_obs_ensemble(book_copy):
    update_workbook(book_copy, obs_csv=OBS_CSV, backend='openpyxl')          # no pst: MODELLED only
    ws = openpyxl.load_workbook(book_copy)['OBS_HEAD']
    hdr = [c.value for c in ws[1]]
    assert 'MODELLED' in hdr and 'RESIDUAL' not in hdr
    assert _col(book_copy, 'OBS_HEAD', 'MODELLED')[0] == pytest.approx(2542.59 + 1)
    update_workbook(book_copy, obs_csv=OBS_CSV, real='0', pst=read_pst(PST), backend='openpyxl', groups=['head'])
    assert _col(book_copy, 'OBS_HEAD', 'MODELLED')[0] == pytest.approx(2542.59 + 3)
    assert _col(book_copy, 'OBS_HEAD', 'RESIDUAL')[0] == pytest.approx(-3)
    assert _col(book_copy, 'OBS_HEAD', 'MODELLED')[2] == pytest.approx(2865.475129 + 1)   # headss: untouched


def test_update_cli_ies(book_copy):
    main(['update', book_copy, '--par', PAR_CSV, '--obs_csv', OBS_CSV, '--pst', PST, '--real', 'best',
          '--group', 'sy,head', '--backend', 'openpyxl'])
    assert _col(book_copy, 'PAR_SY', 'PARVAL1')[0] == pytest.approx(0.15 * 0.5)
    assert _col(book_copy, 'PAR_HK', 'PARVAL1')[0] == pytest.approx(49.73775)
    assert _col(book_copy, 'OBS_HEAD', 'RESIDUAL')[0] == pytest.approx(-0.5)
    with pytest.raises(ValueError, match='not both'):
        main(['update', book_copy, '--res', os.path.join(DATA, 'demo.res'), '--obs_csv', OBS_CSV,
              '--backend', 'openpyxl'])


# ---------------------------------------------------------------------- parrep
def test_parrep_from_par_and_set(tmp_path):
    out = tmp_path / 'rep.pst'
    main(['parrep', PST, os.path.join(DATA, 'demo.par'), str(out), '--set', 'noptmax=0', '--set', 'pestmode=estimation'])
    a = read_pst(PST)
    b = read_pst(out)
    assert np.allclose(b.par['PARVAL1'].astype(float), a.par['PARVAL1'].astype(float) * 2)
    assert b.control['noptmax'] == 0 and b.pestmode == 'estimation'
    assert b.control['jacupdate'] == 999 and b.pestpp == a.pestpp          # everything else kept
    assert list(b.obs['OBSNME']) == list(a.obs['OBSNME'])
    assert not (tmp_path / 'dump.tpl').exists()


def test_parrep_from_ensemble_best(tmp_path):
    out = tmp_path / 'best.pst'
    main(['parrep', PST, PAR_CSV, str(out), '--real', 'best'])
    a = read_pst(PST).par['PARVAL1'].astype(float)
    assert np.allclose(read_pst(out).par['PARVAL1'].astype(float), a * 0.5)
    with pytest.raises(SystemExit):
        main(['parrep', PST, PAR_CSV, str(out), '--set', 'noptmax'])


def test_parrep_warns_out_of_bounds(tmp_path):
    out = tmp_path / 'oob.pst'
    with pytest.warns(UserWarning, match='outside their bounds'):
        main(['parrep', PST, PAR_CSV, str(out), '--real', '0'])      # x3 pushes hk above PARUBND
    assert read_pst(out).npar == 11                                   # still written; PEST will complain


def test_parrep_from_workbook_build_sheet(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)                                   # BUILD names demo.xlsx by bare file name
    out = tmp_path / 'from_book.pst'
    main(['parrep', BOOK, os.path.join(DATA, 'demo.par'), str(out), '--set', 'noptmax=0'])
    ref = tmp_path / 'from_pst.pst'
    main(['parrep', PST, os.path.join(DATA, 'demo.par'), str(ref), '--set', 'noptmax=0'])
    assert out.read_text() == ref.read_text()
    assert not (tmp_path / 'dump.tpl').exists()


def test_parrep_from_dumped_workbook(tmp_path, monkeypatch):
    book = tmp_path / 'sub' / 'dumped.xlsx'
    book.parent.mkdir()
    to_workbook(read_pst(PST), str(book))
    monkeypatch.chdir(tmp_path)                                   # not the workbook's folder
    out = tmp_path / 'rep.pst'
    main(['parrep', str(book), PAR_CSV, str(out), '--real', '1'])
    a = read_pst(PST).par['PARVAL1'].astype(float)
    assert np.allclose(read_pst(out).par['PARVAL1'].astype(float), a * 0.5)


def test_parrep_workbook_without_build_sheet(tmp_path):
    book = tmp_path / 'nobuild.xlsx'
    pd.DataFrame({'PARNME': ['x']}).to_excel(book, sheet_name='PAR', index=False)
    with pytest.raises(SystemExit, match='no BUILD sheet'):
        main(['parrep', str(book), os.path.join(DATA, 'demo.par'), str(tmp_path / 'x.pst')])


# ---------------------------------------------------------------------- sheet globs, build from BUILD sheet
def test_build_with_sheet_globs(tmp_path):
    out = tmp_path / 'glob.pst'
    args = [a for a in build_args(out) if not a.endswith((',PAR_HK', ',PAR_SY', ',PAR_RCH', ',OBS_HEAD', ',OBS_FLOW'))]
    args = [a for i, a in enumerate(args) if not (a in ('--add_par_xls', '--add_obs_xls'))]
    args += ['--add_par_xls', f'{BOOK},PAR_*', '--add_obs_xls', f'{BOOK},obs_*']   # case-insensitive
    main(args + ['--no_manifest'])
    assert out.read_text() == golden()
    with pytest.raises(ValueError, match='no sheet'):
        main(build_args(tmp_path / 'x.pst') + ['--add_par_xls', f'{BOOK},NOPE_*', '--no_manifest'])


def test_build_from_workbook_build_sheet(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    main(['build', BOOK, '--out', str(tmp_path / 'fromsheet.pst'), '--no_manifest'])
    assert (tmp_path / 'fromsheet.pst').read_text() == golden()
    assert not (tmp_path / 'dump.tpl').exists()                    # BUILD sheet says --no_dump_tpl
    shutil.copy(BOOK, tmp_path / 'demo.xlsx')
    main(['build', str(tmp_path / 'demo.xlsx'), '--no_manifest'])   # default: the name in the BUILD sheet, beside the workbook
    assert (tmp_path / 'demo.pst').read_text() == golden()


# ---------------------------------------------------------------------- array formulas, unmatched names
def test_update_leaves_array_formula_cells_alone(book_copy, capsys):
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
    wb = openpyxl.load_workbook(book_copy)
    ws = wb['PAR_SY']
    ws['D2'] = ArrayFormula('D2:D3', '=IF({1;1}, 0.11, 0)')                   # PARVAL1 of sy1_cc01 / sy1_cc02
    wb.save(book_copy)
    upd = update_workbook(book_copy, par=os.path.join(DATA, 'demo.par'), backend='openpyxl')
    assert upd['sheets']['PAR_SY'] == {'rows': 2, 'formula_cells_skipped': 1, 'cells_refused': 0}
    assert 'PAR_SY: 2 rows matched, 1 formula cells left alone' in capsys.readouterr().out
    assert _col(book_copy, 'PAR_HK', 'PARVAL1')[0] == pytest.approx(49.73775 * 2)   # the rest still updated
    upd = update_workbook(book_copy, par=os.path.join(DATA, 'demo.par'), backend='openpyxl', overwrite_formulas=True)
    assert upd['sheets']['PAR_SY']['cells_refused'] == 1                        # asked to, but Excel would refuse


def test_update_reports_unmatched_names(book_copy):
    par = pd.DataFrame({'PARNME': ['hk1_cc01', 'ghost_a', 'ghost_b'], 'PARVAL1': [1.0, 2.0, 3.0]})
    with pytest.warns(UserWarning) as rec:
        upd = update_workbook(book_copy, par=par, backend='openpyxl')
    texts = [str(w.message) for w in rec]
    assert any('2 parameters in the results have no row in the workbook' in t and 'ghost_a' in t for t in texts)
    assert any('10 parameters in the workbook are not in the results' in t for t in texts)
    assert upd['unmatched'] == {'parameters': {'not_in_workbook': 2, 'not_in_results': 10}}
    assert _col(book_copy, 'PAR_HK', 'PARVAL1')[0] == 1.0
